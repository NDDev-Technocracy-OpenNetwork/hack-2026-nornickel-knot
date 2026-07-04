"""Единый пайплайн PDF/DOCX/PPTX/XLS -> чистые, помеченные чанки.

Строится с расчётом на то, чтобы БЫТЬ единственным источником текста -- и
для контракта Данилы (Neo4j), и для последующего LLM-извлечения фактов.
Раньше это были два разных пути (старый плоский парсер для фактов, новый
pymupdf4llm для чанков), из-за чего числа из таблиц терялись в фактах, а
chunk_id между двумя нарезками не совпадал по смыслу. Здесь -- один проход,
один canonical document_id/chunk_id, дальше используется везде.

Правила, встроенные с первого прохода (не патчи после):
  - фикс кириллицы cp1251-как-latin1 (подход Айтала, pdf_to_md.py)
  - чистка нечитаемых символов (U+FFFD)
  - реальные page_start/page_end везде где формат это позволяет
  - честная пометка ocr_noise для мусорных огрызков, не молчаливое "ok"
  - section_title всегда заполнен (перенос заголовка/имя листа/слайда)
  - таблицы (xls, docx-таблицы) идут напрямую как структура, не через OCR

Запуск:
  modal run src/graph/build_corpus_v2.py
"""

import json
import pathlib
import re

import modal

app = modal.App("nornikel-corpus-v2")
ROOT = pathlib.Path.cwd()
YANDEX_ROOT = ROOT / "yandex" / "Задача 2. Научный клубок"
UNPACKED_ROOT = ROOT / "data" / "unpacked_archives"
DOC_CONVERTED_ROOT = ROOT / "data" / "doc_converted"
OUT_PATH = ROOT / "data" / "corpus_v2_chunks.jsonl"
FAILED_PATH = ROOT / "data" / "corpus_v2_failed.jsonl"

CHUNK_CHAR_TARGET = 2500
MIN_REAL_CHARS = 25

image = (
    modal.Image.debian_slim(python_version="3.11")
    .apt_install("tesseract-ocr", "tesseract-ocr-rus", "tesseract-ocr-eng")
    .pip_install(
        "pymupdf4llm==1.28.0", "python-docx==1.1.2", "python-pptx==1.0.2",
        "openpyxl==3.1.5", "xlrd==2.0.1",
    )
    .add_local_dir(str(YANDEX_ROOT), "/root/yandex")
    .add_local_dir(str(UNPACKED_ROOT), "/root/unpacked")
    .add_local_dir(str(DOC_CONVERTED_ROOT), "/root/doc_converted")
)

# PDF -- отдельный образ и GPU: EasyOCR вместо Tesseract, ~6.4x быстрее и
# чище на кириллице (замерено -- src/graph/benchmark_ocr.py, 81с vs 12.6с
# на 128-страничном документе, без путаницы кодировок в выходном тексте).
# Структуру (таблицы/заголовки) для страниц с текстовым слоем всё равно
# даёт pymupdf4llm -- OCR подключается ТОЛЬКО для страниц без текста.
pdf_image = (
    modal.Image.debian_slim(python_version="3.11")
    .pip_install("pymupdf4llm==1.28.0", "easyocr==1.7.2")
    .add_local_dir(str(YANDEX_ROOT), "/root/yandex")
    .add_local_dir(str(UNPACKED_ROOT), "/root/unpacked")
)

# --- Фикс кириллицы и мусора (подход Айтала, применяется универсально ко
# всем форматам, не только PDF) -------------------------------------------

_GARBLED_CP1251_RE = re.compile(r"[À-ÿ]+")


def fix_garbled_cyrillic(text: str) -> str:
    def repl(m: re.Match) -> str:
        chunk = m.group(0)
        try:
            return chunk.encode("latin1").decode("cp1251")
        except (UnicodeEncodeError, UnicodeDecodeError):
            return chunk
    return _GARBLED_CP1251_RE.sub(repl, text)


def clean_unrecoverable_chars(text: str) -> str:
    lines = text.split("\n")
    out = []
    for line in lines:
        if not line:
            out.append(line)
            continue
        ratio = line.count("�") / max(len(line), 1)
        if ratio > 0.3:
            continue  # строка почти целиком нечитаемый мусор -- выбрасываем
        out.append(line.replace("�", "-"))
    return "\n".join(out)


def clean_text(text: str) -> str:
    return clean_unrecoverable_chars(fix_garbled_cyrillic(text))


_NOISE_STRIP_RE = re.compile(r"<!--.*?-->|<br\s*/?>|\\|[|/@#*_~`]", re.DOTALL)


def real_content_len(text: str) -> int:
    stripped = _NOISE_STRIP_RE.sub("", text)
    return len(re.sub(r"\s+", "", stripped))


def slugify(name: str) -> str:
    s = re.sub(r"[^a-zA-Zа-яА-ЯёЁ0-9]+", "_", name.strip().lower())
    return re.sub(r"_+", "_", s).strip("_")[:80]


def classify_chunk(text: str, real_len: int, is_table: bool = False) -> str:
    if is_table:
        return "table"
    table_rows = len(re.findall(r"^\s*\|.*\|\s*$", text, re.MULTILINE))
    if table_rows >= 2:
        return "table"
    if "Start of picture text" in text:
        return "figure_caption" if real_len < 150 else "slide_text"
    low = text.lower()
    if "вывод" in low or "заключен" in low:
        return "conclusion"
    if "эксперимент" in low or "результат" in low:
        return "experiment"
    return "text"


def split_into_chunks(full_text: str, target: int = CHUNK_CHAR_TARGET):
    paragraphs = [p for p in full_text.split("\n\n") if p.strip()]
    chunks, buf, buf_len, consumed = [], [], 0, ""
    for p in paragraphs:
        buf.append(p)
        buf_len += len(p)
        consumed += p + "\n\n"
        if buf_len >= target:
            chunks.append(("\n\n".join(buf), _last_heading(consumed)))
            buf, buf_len = [], 0
    if buf:
        chunks.append(("\n\n".join(buf), _last_heading(consumed)))
    return chunks


def _last_heading(text_before: str):
    matches = re.findall(r"^#{1,6}\s+(.+)$", text_before, re.MULTILINE)
    return matches[-1].strip("* ").strip() if matches else None


def make_rows(document_id: str, source_filename: str, source_format: str,
              sections: list[tuple]) -> list[dict]:
    """sections: список (page_start, page_end, section_title, text, is_table)."""
    rows, idx, last_section = [], 0, None
    for page_start, page_end, section_hint, text, is_table in sections:
        text = clean_text(text).strip()
        if not text:
            continue
        for chunk_text, header_found in split_into_chunks(text):
            section_title = header_found or section_hint or last_section or f"page_{page_start}"
            if header_found:
                last_section = header_found
            real_len = real_content_len(chunk_text)
            rows.append({
                "document_id": document_id, "source_filename": source_filename,
                "source_format": source_format,
                "chunk_id": f"doc_{slugify(document_id)}_chunk_{idx:04d}",
                "chunk_index": idx,
                "chunk_type": classify_chunk(chunk_text, real_len, is_table),
                "section_title": section_title,
                "page_start": page_start, "page_end": page_end,
                "text": chunk_text, "token_count": len(chunk_text) // 4,
                "char_count": len(chunk_text), "real_char_count": real_len,
                "language": "ru",
                "extraction_quality": "ocr_noise" if real_len < MIN_REAL_CHARS else "ok",
            })
            idx += 1
    return rows


# --- Экстракторы по форматам -----------------------------------------------

_easyocr_reader = None


def _get_easyocr_reader():
    global _easyocr_reader
    if _easyocr_reader is None:
        import easyocr
        _easyocr_reader = easyocr.Reader(["ru", "en"], gpu=True)
    return _easyocr_reader


def _process_one_pdf(path: pathlib.Path) -> dict:
    import fitz
    import pymupdf4llm

    stem = path.stem
    try:
        doc = fitz.open(str(path))
    except Exception as e:  # noqa: BLE001
        return {"document_id": stem, "rows": [], "error": str(e)[:300]}

    reader = None
    sections = []
    for i in range(doc.page_count):
        page_no = i + 1
        text_layer = doc[i].get_text("text")
        if len(text_layer.strip()) >= 20:
            # есть текстовый слой -- pymupdf4llm даёт структуру (таблицы,
            # заголовки) без всякого OCR для этой конкретной страницы
            try:
                page_md = pymupdf4llm.to_markdown(doc, pages=[i], show_progress=False)
            except Exception:  # noqa: BLE001
                page_md = text_layer
            sections.append((page_no, page_no, None, page_md, False))
        else:
            # страница без текста -- рендерим в картинку, читаем EasyOCR (GPU)
            if reader is None:
                reader = _get_easyocr_reader()
            pix = doc[i].get_pixmap(dpi=200)
            ocr_lines = reader.readtext(pix.tobytes("png"), detail=0, paragraph=True)
            sections.append((page_no, page_no, None, "\n".join(ocr_lines), False))

    rows = make_rows(stem, path.name, "pdf", sections)
    return {"document_id": stem, "rows": rows, "error": None}


@app.function(image=pdf_image, gpu="H100", timeout=1800, max_containers=20)
def process_pdf(rel_path: str) -> dict:
    path = pathlib.Path("/root") / rel_path
    return _process_one_pdf(path)


@app.function(image=pdf_image, gpu="H100", timeout=7200, max_containers=16)
def process_pdf_batch(rel_paths: list[str], worker_id: int) -> dict:
    """Не документ за раз, а целая группа документов на одном воркере,
    одна загруженная EasyOCR-модель на всю группу -- без пересоздания
    контейнера на каждый файл. Пишет свой JSONL, склейка потом отдельно."""
    out_path = pathlib.Path("/root") / f"worker_{worker_id}.jsonl"
    n_rows, n_failed = 0, 0
    with open(out_path, "w", encoding="utf-8") as f:
        for rel_path in rel_paths:
            path = pathlib.Path("/root") / rel_path
            result = _process_one_pdf(path)
            if result["error"]:
                n_failed += 1
                continue
            for row in result["rows"]:
                f.write(json.dumps(row, ensure_ascii=False) + "\n")
                n_rows += 1
    return {"worker_id": worker_id, "n_docs": len(rel_paths), "n_rows": n_rows,
            "n_failed": n_failed, "content": out_path.read_text(encoding="utf-8")}


@app.function(image=image, timeout=300, max_containers=100)
def process_docx(rel_path: str) -> dict:
    import docx
    path = pathlib.Path("/root") / rel_path
    stem = path.stem
    try:
        d = docx.Document(str(path))
    except Exception as e:  # noqa: BLE001
        return {"document_id": stem, "rows": [], "error": str(e)[:300]}

    sections, buf, current_heading = [], [], None
    for para in d.paragraphs:
        if para.style.name and para.style.name.lower().startswith("heading"):
            if buf:
                sections.append((None, None, current_heading, "\n\n".join(buf), False))
                buf = []
            current_heading = para.text.strip()
            buf.append(f"## {current_heading}")
        elif para.text.strip():
            buf.append(para.text.strip())
    if buf:
        sections.append((None, None, current_heading, "\n\n".join(buf), False))

    for table in d.tables:
        table_lines = []
        for row in table.rows:
            cells = [c.text.strip().replace("\n", " ") for c in row.cells]
            table_lines.append("| " + " | ".join(cells) + " |")
        if len(table_lines) >= 2:
            header_sep = "|" + "---|" * len(table.rows[0].cells)
            table_md = table_lines[0] + "\n" + header_sep + "\n" + "\n".join(table_lines[1:])
            sections.append((None, None, "Таблица", table_md, True))

    rows = make_rows(stem, path.name, "docx", sections)
    return {"document_id": stem, "rows": rows, "error": None}


@app.function(image=image, timeout=300, max_containers=100)
def process_pptx(rel_path: str) -> dict:
    from pptx import Presentation
    path = pathlib.Path("/root") / rel_path
    stem = path.stem
    try:
        prs = Presentation(str(path))
    except Exception as e:  # noqa: BLE001
        return {"document_id": stem, "rows": [], "error": str(e)[:300]}

    sections = []
    for i, slide in enumerate(prs.slides, start=1):
        texts = []
        title = None
        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue
            t = shape.text_frame.text.strip()
            if not t:
                continue
            if title is None and shape == slide.shapes.title:
                title = t
            texts.append(t)
        notes = ""
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
            notes = slide.notes_slide.notes_text_frame.text.strip()
        body = "\n\n".join(texts)
        if notes:
            body += f"\n\n[заметки к слайду]: {notes}"
        sections.append((i, i, title or f"Слайд {i}", body, False))

    rows = make_rows(stem, path.name, "pptx", sections)
    return {"document_id": stem, "rows": rows, "error": None}


@app.function(image=image, timeout=300, max_containers=100)
def process_xls(rel_path: str) -> dict:
    path = pathlib.Path("/root") / rel_path
    stem = path.stem
    sections = []
    try:
        if path.suffix.lower() == ".xlsx":
            import openpyxl
            wb = openpyxl.load_workbook(str(path), data_only=True, read_only=True)
            sheet_names = wb.sheetnames
            get_rows = lambda name: wb[name].iter_rows(values_only=True)  # noqa: E731
        else:
            import xlrd
            wb = xlrd.open_workbook(str(path))
            sheet_names = wb.sheet_names()
            get_rows = lambda name: (  # noqa: E731
                wb.sheet_by_name(name).row_values(r) for r in range(wb.sheet_by_name(name).nrows)
            )
    except Exception as e:  # noqa: BLE001
        return {"document_id": stem, "rows": [], "error": str(e)[:300]}

    for sheet_name in sheet_names:
        try:
            raw_rows = list(get_rows(sheet_name))
        except Exception:  # noqa: BLE001
            continue
        raw_rows = [r for r in raw_rows if any(c not in (None, "") for c in r)]
        if not raw_rows:
            continue
        raw_rows = raw_rows[:500]  # предохранитель от гигантских листов
        md_lines = []
        header = raw_rows[0]
        md_lines.append("| " + " | ".join(str(c) if c is not None else "" for c in header) + " |")
        md_lines.append("|" + "---|" * len(header))
        for row in raw_rows[1:]:
            md_lines.append("| " + " | ".join(str(c) if c is not None else "" for c in row) + " |")
        sections.append((None, None, sheet_name, "\n".join(md_lines), True))

    rows = make_rows(stem, path.name, "xls", sections)
    return {"document_id": stem, "rows": rows, "error": None}


@app.local_entrypoint()
def main():
    def read_list(name):
        p = ROOT / "data" / "file_lists" / name
        if not p.exists():
            return []
        lines = [line.strip() for line in p.read_text(encoding="utf-8").splitlines() if line.strip()]
        return lines

    def to_rel(path_str: str) -> str:
        p = pathlib.Path(path_str).resolve()
        if str(p).startswith(str(YANDEX_ROOT)):
            return "yandex/" + str(p.relative_to(YANDEX_ROOT))
        if str(p).startswith(str(UNPACKED_ROOT)):
            return "unpacked/" + str(p.relative_to(UNPACKED_ROOT))
        if str(p).startswith(str(DOC_CONVERTED_ROOT)):
            return "doc_converted/" + str(p.relative_to(DOC_CONVERTED_ROOT))
        raise ValueError(f"путь вне известных корней: {path_str}")

    pdf_paths = [to_rel(p) for p in read_list("all_pdfs.txt")]
    docx_paths = [to_rel(p) for p in read_list("all_docx.txt")]
    pptx_paths = [to_rel(p) for p in read_list("all_pptx.txt")]
    xls_paths = [to_rel(p) for p in read_list("all_xls.txt")]

    # Возобновление: если файл уже частично собран (прошлый прогон упал/
    # завис), пропускаем документы, чьи чанки уже записаны, вместо того
    # чтобы стирать прогресс и начинать с нуля.
    done_ids: set[str] = set()
    if OUT_PATH.exists():
        with open(OUT_PATH, encoding="utf-8") as f:
            for line in f:
                done_ids.add(json.loads(line)["document_id"])
    if OUT_PATH.exists() and FAILED_PATH.exists():
        with open(FAILED_PATH, encoding="utf-8") as f:
            for line in f:
                doc_id = json.loads(line).get("document_id")
                if doc_id:
                    done_ids.add(doc_id)
    if done_ids:
        print(f"Возобновление: {len(done_ids)} документов уже обработано, пропускаю их.")

    def stem_of(rel_path: str) -> str:
        return pathlib.Path(rel_path).stem

    pdf_paths = [p for p in pdf_paths if stem_of(p) not in done_ids]
    docx_paths = [p for p in docx_paths if stem_of(p) not in done_ids]
    pptx_paths = [p for p in pptx_paths if stem_of(p) not in done_ids]
    xls_paths = [p for p in xls_paths if stem_of(p) not in done_ids]

    print(f"К обработке: PDF {len(pdf_paths)}, DOCX {len(docx_paths)}, PPTX {len(pptx_paths)}, XLS {len(xls_paths)}")

    total_rows, total_failed = 0, 0
    out_f = open(OUT_PATH, "a", encoding="utf-8")
    failed_f = open(FAILED_PATH, "a", encoding="utf-8")

    # PDF -- не по одному документу на вызов, а весь список поделён на
    # N_PDF_WORKERS групп, один воркер = одна загрузка EasyOCR-модели на
    # всю свою группу, не на каждый файл. Каждый пишет свой jsonl внутри
    # контейнера и возвращает содержимое одним куском -- склеиваем здесь.
    N_PDF_WORKERS = 16
    if pdf_paths:
        groups = [pdf_paths[i::N_PDF_WORKERS] for i in range(N_PDF_WORKERS)]
        groups = [(g, idx) for idx, g in enumerate(groups) if g]
        print(f"\n--- pdf: {len(pdf_paths)} файлов, {len(groups)} воркеров ---")
        for result in process_pdf_batch.starmap(groups, return_exceptions=True):
            if isinstance(result, Exception):
                total_failed += 1
                failed_f.write(json.dumps({"format": "pdf", "error": str(result)[:300]}, ensure_ascii=False) + "\n")
                continue
            out_f.write(result["content"])
            total_rows += result["n_rows"]
            total_failed += result["n_failed"]
            print(f"  воркер {result['worker_id']}: {result['n_docs']} документов, "
                  f"{result['n_rows']} чанков, {result['n_failed']} провалов")

    jobs = [
        (process_docx, docx_paths, "docx"),
        (process_pptx, pptx_paths, "pptx"),
        (process_xls, xls_paths, "xls"),
    ]
    for fn, paths, label in jobs:
        if not paths:
            continue
        print(f"\n--- {label}: {len(paths)} файлов ---")
        for i, result in enumerate(fn.map(paths, return_exceptions=True)):
            if isinstance(result, Exception):
                total_failed += 1
                failed_f.write(json.dumps({"format": label, "error": str(result)[:300]}, ensure_ascii=False) + "\n")
                continue
            if result.get("error"):
                total_failed += 1
                failed_f.write(json.dumps({
                    "format": label, "document_id": result["document_id"], "error": result["error"],
                }, ensure_ascii=False) + "\n")
                continue
            rows = result["rows"]
            if not rows:
                failed_f.write(json.dumps({
                    "format": label, "document_id": result["document_id"],
                    "error": "0 чанков (пустой документ после чистки)",
                }, ensure_ascii=False) + "\n")
                continue
            for row in rows:
                out_f.write(json.dumps(row, ensure_ascii=False) + "\n")
            total_rows += len(rows)
            if (i + 1) % 50 == 0:
                print(f"  [{i+1}/{len(paths)}] обработано, всего чанков: {total_rows}")

    out_f.close()
    failed_f.close()
    print(f"\nГотово: {total_rows} чанков, {total_failed} документов с проблемами")
    print(f"-> {OUT_PATH.relative_to(ROOT)}")
    print(f"-> провалы: {FAILED_PATH.relative_to(ROOT)}")
